"""Quarterly numbers (requirement 5.8) and the Excel workbook. Plain SQL and arithmetic, no model."""
from __future__ import annotations

import io
import statistics
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.models import AuditLog, Candidate, CandidateVersion, Request, ShortlistItem, User

OPEN = ("new", "needs_info", "shortlisting", "ready_for_review")


def quarter_range(year: int, quarter: int) -> tuple[datetime, datetime]:
    start_month = 3 * (quarter - 1) + 1
    start = datetime(year, start_month, 1, tzinfo=timezone.utc)
    end = datetime(year + 1, 1, 1, tzinfo=timezone.utc) if quarter == 4 else datetime(year, start_month + 3, 1, tzinfo=timezone.utc)
    return start, end


def current_quarter(now: datetime | None = None) -> tuple[int, int]:
    now = now or datetime.now(timezone.utc)
    return now.year, (now.month - 1) // 3 + 1


def first_shortlist_times(db: Session, request_ids: list[int]) -> dict[int, datetime]:
    rows = db.execute(
        select(AuditLog.entity_id, func.min(AuditLog.ts)).where(
            AuditLog.action == "request.shortlist_ready", AuditLog.entity_type == "request",
            AuditLog.entity_id.in_(request_ids or [-1])).group_by(AuditLog.entity_id)).all()
    return {rid: ts for rid, ts in rows}


def quarterly_report(db: Session, year: int, quarter: int) -> dict[str, Any]:
    start, end = quarter_range(year, quarter)
    reqs = list(db.scalars(select(Request).where(Request.received_at >= start, Request.received_at < end)
                           .order_by(Request.received_at)))
    ids = [r.id for r in reqs]
    owners = {u.id: u.name for u in db.scalars(select(User))}
    items = list(db.scalars(select(ShortlistItem).where(ShortlistItem.request_id.in_(ids or [-1]),
                                                        ShortlistItem.status != "removed")))
    first = first_shortlist_times(db, ids)
    hours = {rid: (first[rid] - r.received_at).total_seconds() / 3600
             for r in reqs for rid in [r.id] if rid in first}

    by_company: dict[str, int] = {}
    by_owner: dict[str, int] = {}
    for r in reqs:
        by_company[r.company or "Unknown"] = by_company.get(r.company or "Unknown", 0) + 1
        who = owners.get(r.owner_id, "Unassigned")
        by_owner[who] = by_owner.get(who, 0) + 1

    per_request = []
    for r in reqs:
        mine = [i for i in items if i.request_id == r.id]
        per_request.append({
            "reference": r.reference, "company": r.company, "job_title": r.job_title, "received": r.received_at,
            "owner": owners.get(r.owner_id, "Unassigned"), "status": r.status,
            "shortlisted": len(mine), "shared": sum(1 for i in mine if i.status == "shared"),
            "hours_to_shortlist": round(hours[r.id], 2) if r.id in hours else None,
        })
    return {
        "year": year, "quarter": quarter, "start": start, "end": end,
        "received": len(reqs), "open": sum(1 for r in reqs if r.status in OPEN),
        "shared_requests": sum(1 for r in reqs if r.status == "shared"), "closed": sum(1 for r in reqs if r.status == "closed"),
        "shortlisted": len({i.candidate_id for i in items}), "shared": len({i.candidate_id for i in items if i.status == "shared"}),
        "median_hours_to_shortlist": round(statistics.median(hours.values()), 2) if hours else None,
        "by_company": sorted(by_company.items(), key=lambda kv: -kv[1]),
        "by_owner": sorted(by_owner.items(), key=lambda kv: -kv[1]),
        "requests": per_request,
    }


def shared_rows(db: Session, request_ids: list[int]) -> list[dict[str, Any]]:
    rows = db.execute(
        select(Request.reference, Request.company, CandidateVersion.full_name, ShortlistItem.decided_at, User.name)
        .join(ShortlistItem, ShortlistItem.request_id == Request.id)
        .join(Candidate, Candidate.id == ShortlistItem.candidate_id)
        .join(CandidateVersion, CandidateVersion.id == Candidate.current_version_id)
        .outerjoin(User, User.id == ShortlistItem.decided_by)
        .where(Request.id.in_(request_ids or [-1]), ShortlistItem.status == "shared")
        .order_by(Request.reference, ShortlistItem.rank)).all()
    return [{"reference": a, "company": b, "candidate": c, "shared_at": d, "by": e} for a, b, c, d, e in rows]


def _naive(value: Any) -> Any:
    return value.replace(tzinfo=None) if isinstance(value, datetime) else value


def _sheet(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_naive(v) for v in row])
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None] or [10])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
    ws.freeze_panes = "A2"


def quarterly_workbook(db: Session, report: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    label = f"Q{report['quarter']} {report['year']}"
    _sheet(ws, ["Measure", label], [
        ["Requests received", report["received"]], ["Still open", report["open"]],
        ["Shared with the company", report["shared_requests"]], ["Closed", report["closed"]],
        ["Candidates shortlisted", report["shortlisted"]], ["Candidates shared", report["shared"]],
        ["Median hours to shortlist", report["median_hours_to_shortlist"]],
    ])
    ws.append([])
    ws.append(["By company", "Requests"])
    for name, n in report["by_company"]:
        ws.append([name, n])
    ws.append([])
    ws.append(["By owner", "Requests"])
    for name, n in report["by_owner"]:
        ws.append([name, n])

    _sheet(wb.create_sheet("Requests"),
           ["Reference", "Company", "Job title", "Received", "Owner", "Status", "Shortlisted", "Shared", "Hours to shortlist"],
           [[r["reference"], r["company"], r["job_title"], r["received"], r["owner"], r["status"], r["shortlisted"],
             r["shared"], r["hours_to_shortlist"]] for r in report["requests"]])
    ids = [r.id for r in db.scalars(select(Request).where(Request.reference.in_([x["reference"] for x in report["requests"]] or ["-"])))]
    _sheet(wb.create_sheet("Shared candidates"), ["Reference", "Company", "Candidate", "Shared at", "Approved by"],
           [[r["reference"], r["company"], r["candidate"], r["shared_at"], r["by"]] for r in shared_rows(db, ids)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
