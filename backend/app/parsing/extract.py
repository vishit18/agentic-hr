"""Text extraction: PDF text layer or DOCX, with OCR fallback for scans and images."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from PIL import Image

from app.config import settings
from app.parsing.safety import FileCheck
from app.parsing.text import clean_text, is_garbage

Image.MAX_IMAGE_PIXELS = settings.max_image_pixels


class ParseError(Exception):
    """The file passed safety checks but no usable text could be read. Message is shown to HR."""


class OcrUnavailable(ParseError):
    pass


@dataclass
class TextResult:
    text: str
    method: str  # pdf_text | docx | ocr
    pages: int | None = None


def pdf_text(path: Path) -> tuple[str, int]:
    pdf = pdfium.PdfDocument(str(path))
    try:
        parts: list[str] = []
        for page in pdf:
            textpage = page.get_textpage()
            parts.append(textpage.get_text_range())
            textpage.close()
            page.close()
        return "\n\n".join(parts), len(pdf)
    finally:
        pdf.close()


def render_pdf_pages(path: Path, scale: float = 2.0) -> list[Image.Image]:
    pdf = pdfium.PdfDocument(str(path))
    try:
        return [page.render(scale=scale).to_pil() for page in pdf]
    finally:
        pdf.close()


def docx_text(path: Path) -> str:
    doc = Document(str(path))
    lines: list[str] = []
    for block in doc.iter_inner_content():
        if hasattr(block, "rows"):  # table
            for row in block.rows:
                cells: list[str] = []
                for cell in row.cells:
                    if cell.text not in cells:  # merged cells repeat
                        cells.append(cell.text)
                lines.append(" | ".join(c for c in cells if c.strip()))
        else:
            lines.append(block.text)
    return "\n".join(lines)


def ocr_images(images: list[Image.Image]) -> str:
    import pytesseract

    try:
        return "\n\n".join(pytesseract.image_to_string(img, lang="eng", timeout=90) for img in images)
    except pytesseract.TesseractNotFoundError as exc:
        raise OcrUnavailable("OCR engine (Tesseract) is not installed.") from exc
    except RuntimeError as exc:  # pytesseract timeout
        raise ParseError("OCR took too long on this file.") from exc


def extract_text(path: Path, check: FileCheck) -> TextResult:
    """Read a file that already passed safety checks. Raises ParseError with a reason for HR."""
    try:
        if check.kind == "pdf":
            raw, pages = pdf_text(path)
            text = clean_text(raw)
            if not is_garbage(text):
                return TextResult(text, "pdf_text", pages)
            text = clean_text(ocr_images(render_pdf_pages(path)))
            if is_garbage(text):
                raise ParseError("No readable text found, even after OCR (scan may be too faint or not English).")
            return TextResult(text, "ocr", pages)

        if check.kind == "docx":
            text = clean_text(docx_text(path))
            if is_garbage(text):
                raise ParseError("Word document has no readable text.")
            return TextResult(text, "docx")

        if check.kind == "image":
            with Image.open(path) as img:
                img.load()
                text = clean_text(ocr_images([img.convert("RGB")]))
            if is_garbage(text):
                raise ParseError("No readable text found in this image.")
            return TextResult(text, "ocr", 1)
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001 - library errors on malformed files become a reason, not a crash
        raise ParseError(f"File could not be read ({exc.__class__.__name__}).") from exc
    raise ParseError("Unsupported file kind.")
