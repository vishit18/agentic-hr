"""Phase 3: tickets, activity, reports and exports, settings and retention."""
import io
import re
import zipfile
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.agents.triage import Summary
from app.config import settings
from app.core import queue
from app.core.audit import audit
from app.core.auth import verify_password
from app.core.models import (
    AppSetting,
    AuditLog,
    Candidate,
    CandidateEmbedding,
    CandidateIdentifier,
    CandidateVersion,
    File,
    Job,
    Request,
    ShortlistItem,
    Ticket,
    User,
)
from app.llm.ollama import ChatResult, LLMUnavailable
from app.worker import run_one
from app.workflows import retention, tickets as tickets_wf
from app.workflows.reporting import quarter_range
from tests import factories
from tests.test_auth_shell import login
from tests.test_candidates_ui import add_candidate


def token(client, path="/home") -> str:
    return re.search(r'name="csrf-token" content="([^"]+)"', client.get(path, follow_redirects=True).text).group(1)


@pytest.fixture()
def staff(client, users):
    login(client, "recruiter@test.local")
    return client


@pytest.fixture()
def admin(client, users):
    login(client, "hr_admin@test.local")
    return client


@pytest.fixture()
def engineer(client, users):
    login(client, "engineer@test.local")
    return client


# ---------------------------------------------------------------- tickets

def seed_agent_trail(db, cand):
    audit(db, actor_type="agent", actor_name="extractor", action="agent.extract", entity_type="candidate_version",
          entity_id=cand.current_version_id, input="the cv text", model="qwen2.5:7b", prompt_version="extractor-v4",
          latency_ms=1234, output={"note": "found jane@secret.test and +973 3600 1234", "experience": 3})
    audit(db, actor_type="system", action="identity.resolved", entity_type="candidate", entity_id=cand.id,
          output={"evidence": ["Same email: jane@secret.test"]})
    db.commit()


def test_mask_pii() -> None:
    out = tickets_wf.mask_pii({"a": "mail jane@secret.test or call +973 3600 1234", "b": ["x@y.org", 5], "c": None})
    assert out == {"a": "mail [email] or call [phone]", "b": ["[email]", 5], "c": None}


def test_flagging_a_problem_attaches_page_records_and_masked_logs_automatically(db, staff, users) -> None:
    cand = add_candidate(db, "Jane Doe", text="Jane Doe cv text")
    seed_agent_trail(db, cand)
    resp = staff.post("/problems", data={"message": "Her experience is missing", "page_url": f"/candidates/{cand.id}?tab=experience"},
                      headers={"X-CSRF-Token": token(staff), "HX-Request": "true"})
    assert resp.status_code == 200 and "problem #1" in resp.text and "Follow it here" in resp.text
    t = db.scalars(select(Ticket)).one()
    assert t.reporter_id == users["recruiter"].id and t.status == "open"
    assert t.record_refs == {"candidate_id": cand.id}
    actions = {r["action"] for r in t.logs["audit"]}
    assert {"agent.extract", "identity.resolved"} <= actions
    dumped = str(t.logs)
    assert "jane@secret.test" not in dumped and "+973 3600 1234" not in dumped and "[email]" in dumped
    row = next(r for r in t.logs["audit"] if r["action"] == "agent.extract")
    assert row["model"] == "qwen2.5:7b" and row["prompt_version"] == "extractor-v4" and row["input_hash"]
    assert db.scalars(select(Job).where(Job.type == "ticket.triage")).one()
    assert db.scalars(select(AuditLog).where(AuditLog.action == "ticket.created")).one().actor_name == "recruiter@test.local"


def test_flag_validation_csrf_and_open_redirect_safety(db, staff) -> None:
    assert staff.post("/problems", data={"message": "no token"}).status_code == 403
    h = {"X-CSRF-Token": token(staff)}
    assert staff.post("/problems", data={"message": "x"}, headers=h).status_code == 400
    staff.post("/problems", data={"message": "something looks off", "page_url": "https://evil.example/x"}, headers=h)
    assert db.scalars(select(Ticket)).one().page_url is None  # only same-site paths are kept


def test_flag_button_and_dialog_are_on_every_staff_page(staff) -> None:
    for path in ("/needs-you", "/candidates", "/requests", "/activity", "/reports"):
        html = staff.get(path).text
        assert "Flag a problem" in html and 'id="flag-dialog"' in html and 'hx-post="/problems"' in html, path


def test_engineer_sees_tickets_logs_and_can_read_attached_agent_input_only(db, client, users) -> None:
    cand = add_candidate(db, "Jane Doe", text="Jane Doe full cv text for the agent")
    seed_agent_trail(db, cand)
    login(client, "recruiter@test.local")
    client.post("/problems", data={"message": "Wrong years", "page_url": f"/candidates/{cand.id}"}, headers={"X-CSRF-Token": token(client)})
    t = db.scalars(select(Ticket)).one()
    client.post("/logout", headers={"X-CSRF-Token": token(client)})
    login(client, "engineer@test.local")
    login_engineer = client

    listing = login_engineer.get("/tickets").text
    assert "Wrong years" in listing and "Open" in listing
    detail = login_engineer.get(f"/tickets/{t.id}").text
    for needle in ("What the HR user wrote", "agent.extract", "qwen2.5:7b", "extractor-v4", "Raw input (logged)", "[email]"):
        assert needle in detail, needle
    assert "jane@secret.test" not in detail

    extract_id = next(r["id"] for r in t.logs["audit"] if r["action"] == "agent.extract")
    shown = login_engineer.get(f"/tickets/{t.id}/input/{extract_id}")
    assert shown.status_code == 200 and "Jane Doe full cv text for the agent" in shown.text
    assert db.scalars(select(AuditLog).where(AuditLog.action == "agent_input.viewed")).one().output["audit_id"] == extract_id

    unrelated = audit(db, actor_type="agent", action="agent.extract", entity_type="candidate_version", entity_id=cand.current_version_id)
    db.commit()
    assert login_engineer.get(f"/tickets/{t.id}/input/{unrelated.id}").status_code == 404  # not attached to this ticket
    assert login_engineer.get("/tickets/99999").status_code == 404


def test_engineer_updates_status_and_reporter_sees_the_fix_note(db, client, users) -> None:
    t = Ticket(reporter_id=users["recruiter"].id, message="Years look wrong", record_refs={}, logs={"audit": [], "jobs": []})
    db.add(t)
    db.commit()
    login(client, "engineer@test.local")
    h = {"X-CSRF-Token": token(client, "/tickets")}
    assert client.post(f"/tickets/{t.id}", data={"status": "nonsense"}, headers=h).status_code == 400
    assert client.post(f"/tickets/{t.id}", data={"status": "fixed", "fix_note": "Dates parser fixed in v2"}).status_code == 403  # csrf
    client.post(f"/tickets/{t.id}", data={"status": "fixed", "fix_note": "Dates parser fixed in v2"}, headers=h)
    db.expire_all()
    assert db.get(Ticket, t.id).status == "fixed"
    client.post("/logout", headers=h)
    login(client, "recruiter@test.local")
    mine = client.get("/problems").text
    assert "Years look wrong" in mine and "Dates parser fixed in v2" in mine and "Fixed" in mine


def test_a_reporter_only_sees_their_own_problems(db, client, users) -> None:
    db.add(Ticket(reporter_id=users["hr_admin"].id, message="Admin's private report", record_refs={}, logs={}))
    db.commit()
    login(client, "recruiter@test.local")
    assert "Admin's private report" not in client.get("/problems").text


def test_triage_agent_summarises_and_falls_back(db, users, session_factory, monkeypatch) -> None:
    class Fake:
        def chat(self, model, system, user, schema):
            return ChatResult(Summary(summary="Experience missing on a candidate profile."), model, 3)

    monkeypatch.setattr(tickets_wf, "get_client", lambda: Fake())
    t = tickets_wf.create_ticket(db, users["recruiter"], "Her experience is missing", "/candidates/1")
    db.commit()
    assert run_one(session_factory, "w")
    db.expire_all()
    assert db.get(Ticket, t.id).summary == "Experience missing on a candidate profile."

    class Down:
        def chat(self, *a, **k):
            raise LLMUnavailable("down")

    monkeypatch.setattr(tickets_wf, "get_client", lambda: Down())
    t2 = tickets_wf.create_ticket(db, users["recruiter"], "x" * 300, None)
    db.commit()
    run_one(session_factory, "w")
    db.expire_all()
    assert db.get(Ticket, t2.id).summary.endswith("...") and len(db.get(Ticket, t2.id).summary) <= 110


# ---------------------------------------------------------------- activity

def test_activity_timeline_uses_plain_english_and_filters(db, staff, users) -> None:
    a = add_candidate(db, "Ann One")
    b = add_candidate(db, "Bob Two")
    audit(db, actor_type="system", actor_name="intake", action="version.written", entity_type="candidate", entity_id=a.id)
    audit(db, actor_type="user", actor_id=users["recruiter"].id, actor_name="recruiter@test.local", action="candidate.merged", entity_type="candidate", entity_id=b.id)
    audit(db, actor_type="system", action="job.done")
    audit(db, actor_type="agent", actor_name="matcher", action="agent.matcher", entity_type="request", entity_id=7, model="qwen2.5:7b")
    db.commit()
    html = staff.get("/activity").text
    assert "Profile updated" in html and "Profiles merged" in html and "Candidate explanation written" in html
    assert "Background task finished" not in html  # routine noise hidden by default
    assert "Background task finished" in staff.get("/activity?noise=1").text
    only_a = staff.get(f"/activity?candidate={a.id}").text
    assert "Profile updated" in only_a and "Profiles merged" not in only_a
    assert "Profiles merged" in staff.get("/activity?user_q=recruiter@").text
    assert "Candidate explanation written" in staff.get("/activity?agent=matcher").text
    assert "Candidate explanation written" in staff.get("/activity?req=7").text
    assert "Nothing matches" in staff.get("/activity?candidate=999999").text


def test_agent_panel_lists_work_in_plain_english_and_polls(db, staff) -> None:
    f = File(sha256="d" * 64, original_name="jane.pdf", stored_path="/x", size_bytes=1, origin="cvs_dir")
    db.add(f)
    db.flush()
    now = datetime.now(timezone.utc)
    db.add_all([
        Job(type="cv.process", payload={"file_id": f.id}, status="running", locked_at=now - timedelta(seconds=7)),
        Job(type="cv.embed", payload={}, status="queued"),
        Job(type="request.process", payload={}, status="dead", last_error="boom", attempts=3, finished_at=now),
        Job(type="cv.process", payload={"file_id": f.id}, status="done", locked_at=now - timedelta(seconds=30), finished_at=now - timedelta(seconds=20)),
    ])
    db.commit()
    page = staff.get("/activity?tab=agents").text
    assert 'hx-get="/activity/agents"' in page and 'hx-trigger="every 3s"' in page
    part = staff.get("/activity/agents").text
    assert "<html" not in part
    for needle in ("Reading a CV", "jane.pdf", "Working", "Waiting", "Failed", "Done", "7 s"):
        assert needle in part, needle
    job = db.scalars(select(Job).where(Job.status == "running")).one()
    audit(db, actor_type="system", action="file.safety", job_id=job.id)
    db.commit()
    detail = staff.get(f"/activity/jobs/{job.id}").text
    assert "File checked" in detail and staff.get("/activity/jobs/99999").status_code == 404


def test_ai_unavailable_banner_for_staff_not_engineers(db, client, users, monkeypatch) -> None:
    from app.api import activity
    from app.core import web

    monkeypatch.setattr(web, "ai_available", lambda: False)
    monkeypatch.setattr(activity, "ai_available", lambda: False)
    login(client, "recruiter@test.local")
    for path in ("/needs-you", "/candidates", "/activity?tab=agents"):
        assert "AI is unavailable" in client.get(path).text or "AI unavailable" in client.get(path).text, path
    client.post("/logout", headers={"X-CSRF-Token": token(client)})
    login(client, "engineer@test.local")
    assert "AI is unavailable" not in client.get("/tickets").text


# ---------------------------------------------------------------- reports and exports

def test_quarter_ranges() -> None:
    assert quarter_range(2026, 1)[0] == datetime(2026, 1, 1, tzinfo=timezone.utc)
    assert quarter_range(2026, 4)[1] == datetime(2027, 1, 1, tzinfo=timezone.utc)
    assert quarter_range(2026, 2) == (datetime(2026, 4, 1, tzinfo=timezone.utc), datetime(2026, 7, 1, tzinfo=timezone.utc))


def seed_quarter(db, users):
    when = datetime(2026, 5, 10, 9, 0, tzinfo=timezone.utc)
    cands = [add_candidate(db, f"Person {i}") for i in range(3)]
    r1 = Request(reference="REQ-2026-0001", company="Globex", job_title="Accountant", status="shared", owner_id=users["recruiter"].id,
                 received_at=when, data={})
    r2 = Request(reference="REQ-2026-0002", company="Globex", job_title="Analyst", status="ready_for_review", owner_id=users["hr_admin"].id,
                 received_at=when + timedelta(days=1), data={})
    r3 = Request(reference="REQ-2026-0003", company="Initech", job_title="Chef", status="closed", received_at=when + timedelta(days=2), data={})
    old = Request(reference="REQ-2025-0009", company="Old Co", status="new", received_at=when - timedelta(days=400), data={})
    db.add_all([r1, r2, r3, old])
    db.flush()
    db.add_all([
        ShortlistItem(request_id=r1.id, candidate_id=cands[0].id, rank=1, status="shared", decided_by=users["recruiter"].id, decided_at=when),
        ShortlistItem(request_id=r1.id, candidate_id=cands[1].id, rank=2, status="suggested"),
        ShortlistItem(request_id=r2.id, candidate_id=cands[1].id, rank=1, status="suggested"),
        ShortlistItem(request_id=r2.id, candidate_id=cands[2].id, rank=2, status="removed"),
    ])
    for r, hours in ((r1, 2), (r2, 4)):
        row = audit(db, actor_type="system", action="request.shortlist_ready", entity_type="request", entity_id=r.id)
        row.ts = r.received_at + timedelta(hours=hours)
    db.commit()
    return r1, r2


def test_report_numbers_for_a_quarter(db, staff, users) -> None:
    seed_quarter(db, users)
    html = staff.get("/reports?year=2026&quarter=2").text
    for needle in ("Q2", "Requests received", "Median hours to shortlist", "Globex", "Initech", "Demo" if False else "Test recruiter"):
        assert needle in html, needle
    m = {k: v for k, v in re.findall(r'<div class="tile-n">([^<]+)</div><div class="muted small">([^<]+)</div>', html)}
    assert m["Requests received"] == "3" if False else True
    tiles = dict((label, value) for value, label in re.findall(r'<div class="tile-n">([^<]+)</div><div class="muted small">([^<]+)</div>', html))
    assert tiles["Requests received"] == "3"
    assert tiles["Still open"] == "1" and tiles["Shared with the company"] == "1" and tiles["Closed"] == "1"
    assert tiles["Candidates shortlisted"] == "2"  # removed items do not count
    assert tiles["Candidates shared"] == "1"
    assert tiles["Median hours to shortlist"] == "3.0"
    assert "REQ-2025-0009" not in html  # other quarter
    empty = staff.get("/reports?year=2026&quarter=1").text
    assert "No requests in Q1 2026" in empty
    assert staff.get("/reports?year=abc").status_code in (200, 422)


def test_excel_export_has_the_numbers_and_is_logged(db, staff, users) -> None:
    seed_quarter(db, users)
    resp = staff.get("/reports/export.xlsx?year=2026&quarter=2")
    assert resp.status_code == 200 and "spreadsheetml" in resp.headers["content-type"]
    assert 'filename="cadence-requests-2026-Q2.xlsx"' in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "Requests", "Shared candidates"]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Requests received"] == 3 and summary["Candidates shared"] == 1 and summary["Globex"] == 2
    reqs = list(wb["Requests"].iter_rows(values_only=True))
    assert reqs[0][0] == "Reference" and {r[0] for r in reqs[1:]} == {"REQ-2026-0001", "REQ-2026-0002", "REQ-2026-0003"}
    shared = list(wb["Shared candidates"].iter_rows(values_only=True))
    assert shared[1][0] == "REQ-2026-0001" and shared[1][2] == "Person 0"
    assert db.scalars(select(AuditLog).where(AuditLog.action == "export.report")).one().output["quarter"] == 2


def test_export_all_is_admin_only_complete_and_has_no_secrets(db, client, users, tmp_path) -> None:
    seed_quarter(db, users)
    cv = factories.write_docx(tmp_path / "jane.docx")
    f = File(sha256="e" * 64, original_name="jane.docx", stored_path=str(cv), size_bytes=cv.stat().st_size, origin="cvs_dir")
    db.add(f)
    db.flush()
    cand = db.scalars(select(Candidate)).first()
    db.get(CandidateVersion, cand.current_version_id).file_id = f.id
    db.commit()

    login(client, "recruiter@test.local")
    assert client.get("/exports/all").status_code == 403
    client.post("/logout", headers={"X-CSRF-Token": token(client)})
    login(client, "hr_admin@test.local")
    resp = client.get("/exports/all")
    assert resp.status_code == 200 and resp.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(resp.content))
    names = set(z.namelist())
    assert {"README.txt", "candidates.csv", "candidates.json", "requests.csv", "requests.json", "review_items.csv",
            "tickets.csv", "audit_log.csv"} <= names
    assert f"files/{f.id}_jane.docx" in names
    assert z.read(f"files/{f.id}_jane.docx") == cv.read_bytes()
    assert "Person 0" in z.read("candidates.csv").decode() and "REQ-2026-0001" in z.read("requests.csv").decode()
    everything = b"".join(z.read(n) for n in names if not n.startswith("files/")).decode()
    assert "$argon2" not in everything and "password_hash" not in everything
    row = db.scalars(select(AuditLog).where(AuditLog.action == "export.all")).one()
    assert row.output["candidates"] == 3 and row.output["requests"] == 4 and row.output["files"] == 1


# ---------------------------------------------------------------- settings: users and retention

def test_settings_is_admin_only(db, client, users) -> None:
    login(client, "recruiter@test.local")
    assert client.get("/settings").status_code == 403
    client.post("/logout", headers={"X-CSRF-Token": token(client)})
    login(client, "hr_admin@test.local")
    page = client.get("/settings").text
    assert "How long CVs are kept" in page and "People who can sign in" in page and "kept until you set a period" in page


def test_admin_creates_a_user_who_can_sign_in_with_a_hashed_password(db, admin, client) -> None:
    h = {"X-CSRF-Token": token(admin)}
    data = {"name": "New Person", "email": "New@Cadence.test", "role": "recruiter", "password": "a-long-password-1"}
    assert admin.post("/settings/users", data=data).status_code == 403  # csrf
    admin.post("/settings/users", data=data, headers=h)
    u = db.scalars(select(User).where(User.email == "new@cadence.test")).one()
    assert u.password_hash.startswith("$argon2") and verify_password(u.password_hash, "a-long-password-1")
    assert "error" in admin.post("/settings/users", data=data, headers=h, follow_redirects=False).headers["location"]  # duplicate
    short = admin.post("/settings/users", data={**data, "email": "x@y.test", "password": "short"}, headers=h, follow_redirects=False)
    assert "error" in short.headers["location"] and db.scalar(select(User.id).where(User.email == "x@y.test")) is None
    assert db.scalars(select(AuditLog).where(AuditLog.action == "user.created")).one()


def test_admin_changes_roles_and_cannot_lock_themselves_out(db, admin, users) -> None:
    h = {"X-CSRF-Token": token(admin)}
    admin.post(f"/settings/users/{users['recruiter'].id}", data={"role": "engineer", "active": ""}, headers=h)
    db.expire_all()
    assert db.get(User, users["recruiter"].id).role == "engineer" and db.get(User, users["recruiter"].id).is_active is False
    resp = admin.post(f"/settings/users/{users['hr_admin'].id}", data={"role": "recruiter", "active": "1"}, headers=h, follow_redirects=False)
    assert "own+administrator" in resp.headers["location"]
    db.expire_all()
    assert db.get(User, users["hr_admin"].id).role == "hr_admin"


def old_candidate(db, name, days_old, *, files=()):
    c = add_candidate(db, name, emails=[f"{name.split()[0].lower()}@x.test"], days_ago=days_old)
    return c


def test_retention_is_off_by_default_and_previews_before_applying(db, admin) -> None:
    assert retention.get_days(db) == 0 and retention.due_candidates(db, 0) == []
    old_candidate(db, "Old Person", 800)
    old_candidate(db, "New Person", 10)
    h = {"X-CSRF-Token": token(admin)}
    resp = admin.post("/settings/retention", data={"days": "365"}, headers=h, follow_redirects=False)
    assert resp.headers["location"] == "/settings?preview=365"  # not applied yet
    assert retention.get_days(db) == 0
    page = admin.get("/settings?preview=365").text
    assert "<strong>1</strong> profile would be removed" in page and "I have checked the number above" in page
    admin.post("/settings/retention", data={"days": "365", "confirm": "1"}, headers=h)
    assert retention.get_days(db) == 365
    assert db.scalars(select(AuditLog).where(AuditLog.action == "retention.changed")).one().output == {"from_days": 0, "to_days": 365}
    assert "error" in admin.post("/settings/retention", data={"days": "-5"}, headers=h, follow_redirects=False).headers["location"]


def test_due_candidates_never_includes_recent_merged_away_or_open_shortlist_profiles(db) -> None:
    old = old_candidate(db, "Old One", 900)
    recent = old_candidate(db, "Recent Two", 30)
    on_list = old_candidate(db, "Listed Three", 900)
    merged = old_candidate(db, "Merged Four", 900)
    merged.merged_into_id = recent.id
    req = Request(reference="R-9", company="X", status="ready_for_review", data={})
    db.add(req)
    db.flush()
    db.add(ShortlistItem(request_id=req.id, candidate_id=on_list.id, rank=1, status="suggested"))
    db.commit()
    assert retention.due_candidates(db, 365) == [old.id]
    req.status = "closed"
    db.commit()
    assert sorted(retention.due_candidates(db, 365)) == sorted([old.id, on_list.id])


def test_purge_removes_profiles_and_files_but_leaves_source_originals(db, tmp_path, monkeypatch, session_factory) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    source = tmp_path / "cvs"
    source.mkdir()
    monkeypatch.setattr(settings, "inbox_dir", inbox)
    uploaded = factories.write_docx(inbox / "u.docx")
    original = factories.write_docx(source / "o.docx")
    fu = File(sha256="1" * 64, original_name="u.docx", stored_path=str(uploaded), size_bytes=1, origin="upload")
    fo = File(sha256="2" * 64, original_name="o.docx", stored_path=str(original), size_bytes=1, origin="cvs_dir")
    db.add_all([fu, fo])
    db.flush()
    a = old_candidate(db, "Gone Soon", 900)
    b = old_candidate(db, "Gone Too", 900)
    keep = old_candidate(db, "Stays Here", 5)
    db.get(CandidateVersion, a.current_version_id).file_id = fu.id
    db.get(CandidateVersion, b.current_version_id).file_id = fo.id
    db.add(CandidateEmbedding(candidate_version_id=a.current_version_id, kind="profile", embedding=[1.0] + [0.0] * (settings.embedding_dim - 1), model="m"))
    db.commit()

    gone_ids, keep_id = [a.id, b.id], keep.id
    retention.set_days(db, 365, None)
    db.commit()
    queue.enqueue(db, "retention.purge", {})
    db.commit()
    assert run_one(session_factory, "w")
    db.expire_all()
    assert {c.id for c in db.scalars(select(Candidate))} == {keep_id}
    assert db.scalar(select(CandidateIdentifier.id).where(CandidateIdentifier.candidate_id.in_(gone_ids))) is None
    assert db.scalar(select(CandidateEmbedding.id)) is None and db.scalar(select(File.id)) is None
    assert not uploaded.exists() and original.exists()  # the read-only source folder is never touched
    row = db.scalars(select(AuditLog).where(AuditLog.action == "retention.purge")).one()
    assert row.output["profiles"] == 2 and row.output["files_deleted"] == 1 and row.output["originals_left_in_source"] == 1
    assert "Gone Soon" not in str(row.output)  # the log records counts, not names


def test_purge_job_does_nothing_when_retention_is_off_and_scheduling_is_once_a_day(db, session_factory) -> None:
    old_candidate(db, "Old Person", 900)
    assert retention.maybe_schedule(db) is False  # off
    queue.enqueue(db, "retention.purge", {})
    db.commit()
    run_one(session_factory, "w")
    db.expire_all()
    assert db.scalar(select(Candidate.id)) is not None
    retention.set_days(db, 365, None)
    db.commit()
    assert retention.maybe_schedule(db) is True
    assert retention.maybe_schedule(db) is False  # already queued today
